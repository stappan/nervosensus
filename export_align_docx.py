import json, re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ── Load cell data from data.js ──────────────────────────────────────
with open('js/data.js', 'r', encoding='utf-8') as f:
    content = f.read()

m = re.search(r'const\s+DEFAULT_CELL_TYPES\s*=\s*', content)
start = m.end()
depth = 0
for i in range(start, len(content)):
    if content[i] == '[': depth += 1
    elif content[i] == ']': depth -= 1
    if depth == 0: break
cells = json.loads(content[start:i+1])

# ── Build ID index ───────────────────────────────────────────────────
id_index = {}
for idx, ct in enumerate(cells):
    if ct.get('id'):
        id_index[ct['id']] = idx

# ── Relationship discovery (mirrors getAssertedRelationships) ────────
def get_relationships(cell_idx):
    ct = cells[cell_idx]
    equivalences, subtypeOf, hasSubtypes = [], [], []

    for m in (ct.get('mapsTo') or []):
        if not m.get('label') and not m.get('id'): continue
        tid = id_index.get(m.get('id'), -1) if m.get('id') else next((i for i, c in enumerate(cells) if c.get('entity') == m.get('label')), -1)
        if tid != -1: equivalences.append(tid)

    for oi, other in enumerate(cells):
        if oi == cell_idx: continue
        for m in (other.get('mapsTo') or []):
            matches = (m.get('id') and ct.get('id') and m['id'] == ct['id']) or (m.get('label') == ct.get('entity'))
            if matches and oi not in equivalences:
                equivalences.append(oi)

    for s in (ct.get('assertedSubclassOf') or []):
        if not s.get('label') and not s.get('id'): continue
        tid = id_index.get(s.get('id'), -1) if s.get('id') else next((i for i, c in enumerate(cells) if c.get('entity') == s.get('label')), -1)
        if tid != -1: subtypeOf.append(tid)

    for oi, other in enumerate(cells):
        if oi == cell_idx: continue
        for s in (other.get('assertedSubclassOf') or []):
            matches = (s.get('id') and ct.get('id') and s['id'] == ct['id']) or (s.get('label') == ct.get('entity'))
            if matches: hasSubtypes.append(oi)

    return equivalences + subtypeOf + hasSubtypes

# ── Helpers ──────────────────────────────────────────────────────────
def get_marker_genes(ct):
    genes = ct.get('markerGenes') or []
    return ', '.join(g.get('name', '') for g in genes if g.get('name'))

def get_temperature(ct):
    attrs = ct.get('clusterAttributes') or {}
    cold = attrs.get('cold_sensitive', False)
    heat = attrs.get('heat_sensitive', False)
    if cold and heat: return 'Cold, Heat'
    if cold: return 'Cold'
    if heat: return 'Heat'
    return ''

def bool_mark(ct, key):
    return '✓' if (ct.get('clusterAttributes') or {}).get(key) else '—'

# ── Build align data ─────────────────────────────────────────────────
big_drg = [(idx, ct) for idx, ct in enumerate(cells) if ct.get('sourceNomenclatureLabel') == 'big DRG paper']

parent_rows = []
for parent_idx, parent_ct in big_drg:
    rel_idxs = get_relationships(parent_idx)
    children = []
    seen = set()
    for ri in rel_idxs:
        if ri != parent_idx and cells[ri].get('sourceNomenclatureLabel') != 'big DRG paper' and ri not in seen:
            children.append((ri, cells[ri]))
            seen.add(ri)
    children.sort(key=lambda x: x[1].get('sourceNomenclatureLabel', ''))
    parent_rows.append((parent_idx, parent_ct, children))

with_children = [(pi, pc, ch) for pi, pc, ch in parent_rows if ch]
without_children = [(pi, pc, ch) for pi, pc, ch in parent_rows if not ch]
with_children.sort(key=lambda x: x[1].get('preferredLabel', ''))
without_children.sort(key=lambda x: x[1].get('preferredLabel', ''))
sorted_rows = with_children + without_children

# ── Column definitions ───────────────────────────────────────────────
columns = [
    ('Cell Type', 'name'),
    ('Source', 'source'),
    ('Marker Genes', 'genes'),
    ('Aβ', 'fiber_a_beta'),
    ('Aδ', 'fiber_a_delta'),
    ('C', 'fiber_c'),
    ('LTM', 'mechanosensitive_ltm'),
    ('HTM', 'mechanosensitive_htm'),
    ('Temp', 'temperature'),
    ('Proprio', 'proprioceptive'),
    ('RA', 'rapidly_adapting'),
    ('SA', 'slowly_adapting'),
    ('Species', 'species'),
]

def cell_values(ct):
    return [
        ct.get('localLabel') or ct.get('preferredLabel', ''),
        ct.get('sourceNomenclatureLabel', ''),
        get_marker_genes(ct),
        bool_mark(ct, 'fiber_a_beta'),
        bool_mark(ct, 'fiber_a_delta'),
        bool_mark(ct, 'fiber_c'),
        bool_mark(ct, 'mechanosensitive_ltm'),
        bool_mark(ct, 'mechanosensitive_htm'),
        get_temperature(ct) or '—',
        bool_mark(ct, 'proprioceptive'),
        bool_mark(ct, 'rapidly_adapting'),
        bool_mark(ct, 'slowly_adapting'),
        ct.get('species', ''),
    ]

# ── Create Word document ─────────────────────────────────────────────
doc = Document()

# Landscape orientation
section = doc.sections[0]
section.orientation = 1  # WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

# Title
title_para = doc.add_paragraph()
title_run = title_para.add_run('NervoSensus — Cross-Source Alignment')
title_run.font.size = Pt(16)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x2d, 0x37, 0x48)

# Subtitle
sub_para = doc.add_paragraph()
sub_run = sub_para.add_run('Big DRG paper cell types with related cells from other sources. ')
sub_run.font.size = Pt(9)
sub_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
sub_run2 = sub_para.add_run(f'{len(big_drg)} parent cells · {sum(len(ch) for _, _, ch in sorted_rows)} related cells')
sub_run2.font.size = Pt(9)
sub_run2.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

# Legend
legend_para = doc.add_paragraph()
legend_run = legend_para.add_run('✓ = feature present    — = feature absent or not reported')
legend_run.font.size = Pt(8)
legend_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
legend_run.font.italic = True

# Count total rows for the table
total_data_rows = 0
for pi, pc, children in sorted_rows:
    total_data_rows += 1 + len(children)

# Add section header rows
section_header_count = 1 if with_children else 0
if without_children:
    section_header_count += 1
total_rows = 1 + section_header_count + total_data_rows  # header + section headers + data

# Create table
table = doc.add_table(rows=1, cols=len(columns))
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# ── Helper to shade a cell ───────────────────────────────────────────
def shade_cell(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_text(cell, text, bold=False, font_size=7, color=None, align_center=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.size = Pt(font_size)
    run.font.name = 'Arial'
    if bold: run.font.bold = True
    if color: run.font.color.rgb = color
    # Tight cell margins
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)

# ── Header row ───────────────────────────────────────────────────────
header_row = table.rows[0]
# Repeat header on each page
trPr = header_row._tr.get_or_add_trPr()
tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")} />')
trPr.append(tblHeader)

for ci, (col_name, _) in enumerate(columns):
    cell = header_row.cells[ci]
    shade_cell(cell, '667eea')
    set_cell_text(cell, col_name, bold=True, font_size=7, color=RGBColor(0xFF, 0xFF, 0xFF), align_center=True)

# ── Column widths ────────────────────────────────────────────────────
col_widths = [Cm(7.5), Cm(4.0), Cm(3.5), Cm(1.2), Cm(1.2), Cm(1.0), Cm(1.2), Cm(1.2), Cm(1.8), Cm(1.4), Cm(1.0), Cm(1.0), Cm(1.5)]
for ci, w in enumerate(col_widths):
    for row in table.rows:
        row.cells[ci].width = w

# ── Populate rows ────────────────────────────────────────────────────
def add_section_header(text):
    row = table.add_row()
    # Merge all cells
    for ci in range(1, len(columns)):
        row.cells[0].merge(row.cells[ci])
    shade_cell(row.cells[0], 'e8ecfb')
    set_cell_text(row.cells[0], text, bold=True, font_size=8, color=RGBColor(0x2d, 0x37, 0x48))

def add_data_row(ct, is_parent=False):
    row = table.add_row()
    vals = cell_values(ct)
    for ci, val in enumerate(vals):
        cell = row.cells[ci]
        is_bool_col = columns[ci][1] in ('fiber_a_beta', 'fiber_a_delta', 'fiber_c',
            'mechanosensitive_ltm', 'mechanosensitive_htm', 'proprioceptive',
            'rapidly_adapting', 'slowly_adapting', 'temperature')
        color = None
        if val == '✓': color = RGBColor(0x22, 0xc5, 0x5e)
        elif val == '—': color = RGBColor(0xcb, 0xd5, 0xe1)
        align = is_bool_col or columns[ci][1] == 'species'
        set_cell_text(cell, val, bold=is_parent and ci == 0, font_size=7, color=color, align_center=align)
        if not is_parent:
            shade_cell(cell, 'f8fafc')
    return row

# Section: cells with alignments
if with_children:
    add_section_header(f'Cells with cross-source alignments ({len(with_children)})')

row_idx = 0
for pi, pc, children in sorted_rows:
    if row_idx == len(with_children) and without_children:
        add_section_header(f'Cells without cross-source alignments ({len(without_children)})')
    add_data_row(pc, is_parent=True)
    for ci_idx, (child_idx, child_ct) in enumerate(children):
        add_data_row(child_ct, is_parent=False)
    row_idx += 1

# ── Keep parent + children together (cantSplit) ──────────────────────
# We can't truly do "keep with next" per-group in python-docx easily,
# but we can set cantSplit on each row to prevent a row from splitting
# across pages, and use keepNext on parent rows.
for row in table.rows:
    trPr = row._tr.get_or_add_trPr()
    cant_split = parse_xml(f'<w:cantSplit {nsdecls("w")} />')
    trPr.append(cant_split)

# Set keepNext on parent rows to keep them with their first child
# We track which rows are parents by walking through again
row_list = list(table.rows)
r = 1  # skip header
for pi, pc, children in sorted_rows:
    if r >= len(row_list): break
    # Check if this is a section header (merged cell)
    cell_text = row_list[r].cells[0].text
    if cell_text.startswith('Cells with') or cell_text.startswith('Cells without'):
        r += 1
    if r >= len(row_list): break
    # This should be the parent row — set keepNext if it has children
    if children:
        pPr = row_list[r]._tr.get_or_add_trPr()
        # keepNext doesn't exist on tr, but we can set it on the paragraph
        for cell in row_list[r].cells:
            for p in cell.paragraphs:
                pPr_p = p._p.get_or_add_pPr()
                keepNext = parse_xml(f'<w:keepNext {nsdecls("w")} />')
                pPr_p.append(keepNext)
    r += 1
    # Skip children
    for _ in children:
        # Set keepNext on all children except the last to keep the group together
        if r < len(row_list) and _ != children[-1]:
            for cell in row_list[r].cells:
                for p in cell.paragraphs:
                    pPr_p = p._p.get_or_add_pPr()
                    keepNext = parse_xml(f'<w:keepNext {nsdecls("w")} />')
                    pPr_p.append(keepNext)
        r += 1

# ── Feedback link at the bottom ──────────────────────────────────────
doc.add_paragraph()  # spacer
feedback_para = doc.add_paragraph()
feedback_run = feedback_para.add_run('We value your feedback on NervoSensus. Submit feedback at: ')
feedback_run.font.size = Pt(9)
feedback_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

from docx.oxml import OxmlElement
hyperlink = OxmlElement('w:hyperlink')
hyperlink.set(qn('w:history'), '1')
# Create the rId for the hyperlink
r_id = doc.part.relate_to('https://stappan.github.io/nervosensus/nervosensus-feedback.html?view=Align%20View', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
hyperlink.set(qn('r:id'), r_id)
new_run = OxmlElement('w:r')
rPr = OxmlElement('w:rPr')
rStyle = OxmlElement('w:rStyle')
rStyle.set(qn('w:val'), 'Hyperlink')
rPr.append(rStyle)
sz = OxmlElement('w:sz')
sz.set(qn('w:val'), '18')  # 9pt = 18 half-points
rPr.append(sz)
color_el = OxmlElement('w:color')
color_el.set(qn('w:val'), '667eea')
rPr.append(color_el)
u_el = OxmlElement('w:u')
u_el.set(qn('w:val'), 'single')
rPr.append(u_el)
new_run.append(rPr)
new_run.text = 'nervosensus-feedback.html'
hyperlink.append(new_run)
feedback_para._p.append(hyperlink)

# ── Save ─────────────────────────────────────────────────────────────
out_path = 'NervoSensus_Align.docx'
doc.save(out_path)
total_children = sum(len(ch) for _, _, ch in sorted_rows)
print(f'Saved {out_path}')
print(f'  {len(big_drg)} Big DRG parent cells')
print(f'  {total_children} related child cells')
print(f'  {len(big_drg) + total_children} total data rows')
