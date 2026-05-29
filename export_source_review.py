"""
Export per-source review workbook from NervoSensus data.

Reads the raw NPO spreadsheet (PRECISIONcelltypeNPO.xlsx) and generates an
Excel workbook with one sheet per source. Property values are pulled directly
from the source spreadsheet so authors can verify exactly what was ingested.

Usage:
    python export_source_review.py                              # all sources
    python export_source_review.py --source "Krauter et al., 2025"  # one source
"""
import json, re, sys, argparse
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

# ── Configuration ───────────────────────────────────────────────────
SOURCE_SHORT = {
    'big DRG paper': 'Big DRG',
    'CSA paper': 'CSA',
    'Krauter et al., 2025': 'Krauter 2025',
    'Qi et al., 2024': 'Qi 2024',
    'Yu et al., 2024': 'Yu 2024',
    'Tavares-Ferreira et al., 2022': 'Tavares-F 2022',
    'Kupari et al., 2021': 'Kupari 2021',
}

COLUMNS = [
    ('Status', 16),
    ('Comments', 30),
    ('Local Label', 25),
    ('Species', 10),
    ('Marker Genes', 45),
    ('Fiber Type', 32),
    ('Functional Phenotype', 32),
    ('Threshold Phenotype', 30),
    ('Adaptation', 28),
    ('Cre Line', 32),
    ('npokb ID', 13),
]

HEADER_ROW = 5  # column headers on this row; data starts at HEADER_ROW + 1

# ── Parse arguments ─────────────────────────────────────────────────
parser = argparse.ArgumentParser(description='Export per-source review workbook')
parser.add_argument('--source', default=None, help='Single source to export (default: all)')
args = parser.parse_args()

# ── Load authoritative cell list from data.js ───────────────────────
with open('js/data.js', 'r', encoding='utf-8') as f:
    content = f.read()
m = re.search(r'const\s+DEFAULT_CELL_TYPES\s*=\s*', content)
start = m.end()
depth = 0
for i in range(start, len(content)):
    if content[i] == '[': depth += 1
    elif content[i] == ']': depth -= 1
    if depth == 0: break
cells = json.loads(content[start:i+1])

cell_by_id = {}
for c in cells:
    cid = c.get('id', '')
    if cid:
        cell_by_id[cid] = c

# ── Read raw property data from source spreadsheet ──────────────────
xlsx_path = 'PRECISIONcelltypeNPO.xlsx'
wb_src = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
ws_src = wb_src['PNS Cells to NPO']

raw_data = {}
for row in ws_src.iter_rows(min_row=2, max_col=11):
    npokb_id = str(row[0].value or '').strip()
    prop = str(row[3].value or '').strip()
    val_label = str(row[6].value or '').strip()
    modifier = str(row[8].value or '').strip()
    action = str(row[10].value or '').strip().lower()

    if not npokb_id or not prop or action == "don't add":
        continue
    if npokb_id not in cell_by_id:
        continue
    if npokb_id not in raw_data:
        raw_data[npokb_id] = {}
    if prop not in raw_data[npokb_id]:
        raw_data[npokb_id][prop] = []

    if val_label and not val_label.startswith('http'):
        if prop == 'ilxtr:hasNucleicAcidExpressionPhenotype' and modifier:
            raw_data[npokb_id][prop].append(f'{val_label} [{modifier}]')
        elif val_label not in raw_data[npokb_id][prop]:
            raw_data[npokb_id][prop].append(val_label)

wb_src.close()

# ── Group cells by source ───────────────────────────────────────────
sources = {}
for cid, c in cell_by_id.items():
    src = c.get('sourceNomenclatureLabel', '')
    if not src:
        continue
    if args.source and src != args.source:
        continue
    if src not in sources:
        sources[src] = {'doi': c.get('sourceNomenclature', ''), 'cells': []}
    sources[src]['cells'].append(cid)

for src in sources:
    sources[src]['cells'].sort(
        key=lambda cid: (cell_by_id[cid].get('localLabel') or cell_by_id[cid].get('preferredLabel', '')).lower()
    )

if not sources:
    print(f'No cells found for source: {args.source}')
    sys.exit(1)

# ── Helpers ─────────────────────────────────────────────────────────
def get_raw_values(npokb_id, prop_name):
    return ', '.join(raw_data.get(npokb_id, {}).get(prop_name, []))

# ── Create workbook ─────────────────────────────────────────────────
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Styles
title_font = Font(name='Arial', bold=True, size=14, color='2d3748')
doi_font = Font(name='Arial', size=10, color='667eea', underline='single')
info_font = Font(name='Arial', size=10, color='718096', italic=True)
data_font = Font(name='Arial', size=10)
needs_review_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
reviewed_fill = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
corrected_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')

for src_name in sorted(sources.keys()):
    src_info = sources[src_name]
    sheet_name = SOURCE_SHORT.get(src_name, src_name)[:31]
    ws = wb.create_sheet(title=sheet_name)

    cell_ids = src_info['cells']
    doi_url = src_info['doi']
    doi_text = re.sub(r'^https?://doi\.org/', '', doi_url) if doi_url else ''

    # ── Header rows (source info) ───────────────────────────────────
    num_cols = len(COLUMNS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    ws['A1'] = f'Source: {src_name}'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(vertical='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    ws['A2'] = f'DOI: {doi_text}'
    ws['A2'].font = doi_font
    if doi_url:
        ws['A2'].hyperlink = doi_url

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_cols)
    ws['A3'] = f'{len(cell_ids)} cells  ·  Generated {date.today().isoformat()}  ·  Please review each row and update the Status column'
    ws['A3'].font = info_font

    # Row 4: blank spacer

    # ── Column headers (row 5) ──────────────────────────────────────
    for ci, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=HEADER_ROW, column=ci, value=col_name)
        ws.column_dimensions[get_column_letter(ci)].width = col_width

    # ── Data rows ───────────────────────────────────────────────────
    for ri, cid in enumerate(cell_ids, HEADER_ROW + 1):
        ct = cell_by_id[cid]

        ws.cell(row=ri, column=1, value='Needs Review').font = data_font
        # Column 2 (Comments) left blank
        ws.cell(row=ri, column=3, value=ct.get('localLabel') or ct.get('preferredLabel', '')).font = data_font
        ws.cell(row=ri, column=4, value=ct.get('species', '')).font = data_font
        ws.cell(row=ri, column=5, value=get_raw_values(cid, 'ilxtr:hasNucleicAcidExpressionPhenotype')).font = data_font
        ws.cell(row=ri, column=6, value=get_raw_values(cid, 'ilxtr:hasAxonPhenotype')).font = data_font
        ws.cell(row=ri, column=7, value=get_raw_values(cid, 'ilxtr:hasFunctionalPhenotype')).font = data_font
        ws.cell(row=ri, column=8, value=get_raw_values(cid, 'ilxtr:hasThresholdPhenotype')).font = data_font
        ws.cell(row=ri, column=9, value=get_raw_values(cid, 'ilxtr:hasAdaptationPhenotype')).font = data_font
        ws.cell(row=ri, column=10, value=get_raw_values(cid, 'ilxtr:hasDriverExpressionConstitutivePhenotype')).font = data_font
        ws.cell(row=ri, column=11, value=cid).font = data_font

    last_row = HEADER_ROW + len(cell_ids)

    # ── Excel Table ─────────────────────────────────────────────────
    safe_name = re.sub(r'[^A-Za-z0-9_]', '', sheet_name.replace(' ', '_').replace('.', ''))
    table_ref = f'A{HEADER_ROW}:{get_column_letter(num_cols)}{last_row}'
    tab = Table(displayName=f'Review_{safe_name}', ref=table_ref)
    tab.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium2', showFirstColumn=False,
        showLastColumn=False, showRowStripes=True
    )
    ws.add_table(tab)

    # ── Data validation: Status dropdown ────────────────────────────
    dv = DataValidation(
        type='list',
        formula1='"Needs Review,Reviewed,Corrected"',
        allow_blank=False,
    )
    dv.error = 'Please select: Needs Review, Reviewed, or Corrected'
    dv.errorTitle = 'Invalid Status'
    ws.add_data_validation(dv)
    dv.add(f'A{HEADER_ROW + 1}:A{last_row}')

    # ── Conditional formatting: Status colors ───────────────────────
    status_range = f'A{HEADER_ROW + 1}:A{last_row}'
    ws.conditional_formatting.add(status_range,
        CellIsRule(operator='equal', formula=['"Needs Review"'], fill=needs_review_fill))
    ws.conditional_formatting.add(status_range,
        CellIsRule(operator='equal', formula=['"Reviewed"'], fill=reviewed_fill))
    ws.conditional_formatting.add(status_range,
        CellIsRule(operator='equal', formula=['"Corrected"'], fill=corrected_fill))

    # ── Freeze panes: lock columns A-C and header rows ──────────────
    ws.freeze_panes = f'D{HEADER_ROW + 1}'

    # ── Row height for header rows ──────────────────────────────────
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[4].height = 8  # thin spacer

# ── Save ────────────────────────────────────────────────────────────
out_path = 'NervoSensus_Source_Review.xlsx'
wb.save(out_path)
print(f'Saved {out_path}')
print(f'  {len(sources)} source sheets:')
for src_name in sorted(sources.keys()):
    short = SOURCE_SHORT.get(src_name, src_name)
    print(f'    {short}: {len(sources[src_name]["cells"])} cells')
